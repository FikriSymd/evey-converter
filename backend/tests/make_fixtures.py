"""Generate sample input files for exercising every tool."""

from __future__ import annotations

import sys
from pathlib import Path

import pymupdf

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/evey-fixtures")
OUT.mkdir(parents=True, exist_ok=True)


def make_pdf(path: Path, pages: int = 3, label: str = "evey") -> None:
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page(width=595, height=842)  # A4
        page.insert_text((72, 120), f"{label} - halaman {i + 1}", fontsize=28)
        page.insert_text(
            (72, 170),
            "Baris contoh untuk pengujian ekstraksi teks dan OCR.",
            fontsize=12,
        )
        # A simple table-like block so pdf_excel has something to find.
        for row in range(3):
            y = 240 + row * 24
            page.insert_text((72, y), f"Item {row + 1}", fontsize=11)
            page.insert_text((200, y), str((row + 1) * 100), fontsize=11)
    doc.save(str(path))
    doc.close()


make_pdf(OUT / "laporan.pdf", 3, "Laporan Q3")
make_pdf(OUT / "lampiran.pdf", 2, "Lampiran")

# DOCX
import docx

d = docx.Document()
d.add_heading("Dokumen Uji", level=1)
d.add_paragraph("Paragraf pertama untuk konversi Word ke PDF.")
d.add_paragraph("Paragraf kedua, dengan angka 12345.")
d.save(str(OUT / "dokumen.docx"))

# XLSX
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.append(["Nama", "Jumlah"])
ws.append(["Item A", 10])
ws.append(["Item B", 20])
wb.save(str(OUT / "tabel.xlsx"))

# PPTX
from pptx import Presentation

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "Slide Uji"
prs.save(str(OUT / "presentasi.pptx"))

# Images
from PIL import Image, ImageDraw

for name, color in (("gambar1.png", (200, 60, 40)), ("gambar2.png", (30, 90, 160))):
    im = Image.new("RGB", (800, 600), (255, 255, 255))
    dr = ImageDraw.Draw(im)
    dr.rectangle([40, 40, 760, 560], outline=color, width=8)
    dr.text((80, 80), f"Gambar uji {name}", fill=(0, 0, 0))
    im.save(OUT / name)

# HTML
(OUT / "halaman.html").write_text(
    "<!DOCTYPE html><html><head><meta charset='utf-8'><title>Uji</title></head>"
    "<body><h1>Halaman Uji</h1><p>Konten HTML untuk konversi ke PDF.</p></body></html>",
    encoding="utf-8",
)

print("fixtures in", OUT)
for f in sorted(OUT.iterdir()):
    print(f" {f.name:24} {f.stat().st_size:>8} bytes")
