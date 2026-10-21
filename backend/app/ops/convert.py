"""Convert operations: PDF <-> Office, PDF <-> image, HTML -> PDF, OCR."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

import pymupdf

from .. import toolchain
from ..config import MAX_PAGES, OP_TIMEOUT
from ..util import OpError, guard_pages, safe_stem


def _run(cmd: list[str], what: str) -> subprocess.CompletedProcess:
    try:
        proc = subprocess.run(cmd, capture_output=True, timeout=OP_TIMEOUT)
    except FileNotFoundError as exc:
        raise OpError(f"{what} tidak terpasang di server.") from exc
    except subprocess.TimeoutExpired as exc:
        raise OpError("Proses melebihi batas waktu.") from exc
    if proc.returncode != 0:
        lines = proc.stderr.decode("utf-8", "ignore").strip().splitlines()
        raise OpError(lines[-1] if lines else f"{what} gagal memproses file.")
    return proc


def _soffice() -> str:
    exe = toolchain.find_soffice()
    if not exe:
        raise OpError(
            "LibreOffice belum terpasang di komputer ini. "
            "Lihat panduan setup, bagian memasang program pendukung."
        )
    return exe


def _soffice_convert(src: Path, out_dir: Path, target: str) -> Path:
    """LibreOffice headless convert. A private profile keeps parallel jobs apart."""
    profile = Path(tempfile.mkdtemp(prefix="lo-profile-"))
    try:
        _run(
            [
                _soffice(),
                "--headless",
                "--norestore",
                "--nolockcheck",
                f"-env:UserInstallation={Path(profile).as_uri()}",
                "--convert-to",
                target,
                "--outdir",
                str(out_dir),
                str(src),
            ],
            "LibreOffice",
        )
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    produced = out_dir / f"{src.stem}.{target.split(':')[0]}"
    if not produced.exists():
        raise OpError("LibreOffice tidak menghasilkan file.")
    clean = out_dir / f"{safe_stem(src.name)}.{target.split(':')[0]}"
    if clean != produced:
        produced.replace(clean)
    return clean


def _open_pdf(path: Path) -> "pymupdf.Document":
    """Open a PDF and refuse it if it exceeds the page ceiling."""
    try:
        doc = pymupdf.open(str(path))
    except Exception as exc:
        raise OpError(f"PDF '{path.name}' tidak bisa dibuka.") from exc
    guard_pages(doc, path.name, MAX_PAGES)
    return doc


def _text_docx(pdf: Path, out: Path) -> None:
    """Fallback PDF -> Word: put extracted text into a real .docx."""
    import docx

    document = docx.Document()
    with _open_pdf(pdf) as doc:
        for page in doc:
            for block in page.get_text("blocks"):
                text = block[4].strip()
                if text:
                    document.add_paragraph(text)
    document.save(str(out))


def pdf_word(inputs: list[Path], opts: dict, out_dir: Path) -> list[Path]:
    src = inputs[0]
    out = out_dir / f"{safe_stem(src.name)}.docx"
    # Page ceiling first: pdf2docx would otherwise chew a huge file before we
    # ever get the chance to refuse it.
    with _open_pdf(src):
        pass
    try:
        from pdf2docx import Converter

        cv = Converter(str(src))
        try:
            cv.convert(str(out))
        finally:
            cv.close()
        if out.exists() and out.stat().st_size > 0:
            return [out]
    except Exception:
        pass
    _text_docx(src, out)
    return [out]


def pdf_excel(inputs: list[Path], opts: dict, out_dir: Path) -> list[Path]:
    import pdfplumber
    from openpyxl import Workbook

    src = inputs[0]
    out = out_dir / f"{safe_stem(src.name)}.xlsx"
    with _open_pdf(src):
        pass
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    found_tables = False
    with pdfplumber.open(str(src)) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                found_tables = True
                for row in table:
                    ws.append(["" if c is None else str(c) for c in row])
                ws.append([])
            if not found_tables:
                text = page.extract_text() or ""
                for line in text.splitlines():
                    ws.append([line])
    wb.save(str(out))
    return [out]


def pdf_ppt(inputs: list[Path], opts: dict, out_dir: Path) -> list[Path]:
    from pptx import Presentation
    from pptx.util import Emu

    src = inputs[0]
    out = out_dir / f"{safe_stem(src.name)}.pptx"
    prs = Presentation()
    with _open_pdf(src) as doc:
        for page in doc:
            rect = page.rect
            prs.slide_width = Emu(int(rect.width * 9525))
            prs.slide_height = Emu(int(rect.height * 9525))
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            pix = page.get_pixmap(dpi=120)
            img = out_dir / f"_pg{page.number + 1}.png"
            pix.save(str(img))
            slide.shapes.add_picture(
                str(img), 0, 0, width=prs.slide_width, height=prs.slide_height
            )
            img.unlink(missing_ok=True)
    prs.save(str(out))
    return [out]


def word_pdf(inputs: list[Path], opts: dict, out_dir: Path) -> list[Path]:
    return [_soffice_convert(inputs[0], out_dir, "pdf")]


def excel_pdf(inputs: list[Path], opts: dict, out_dir: Path) -> list[Path]:
    return [_soffice_convert(inputs[0], out_dir, "pdf")]


def ppt_pdf(inputs: list[Path], opts: dict, out_dir: Path) -> list[Path]:
    return [_soffice_convert(inputs[0], out_dir, "pdf")]


def pdf_jpg(inputs: list[Path], opts: dict, out_dir: Path) -> list[Path]:
    src = inputs[0]
    fmt = str(opts.get("imgfmt", "jpg")).lower()
    dpi = int(str(opts.get("dpi", "150")))
    ext = "png" if fmt == "png" else "jpg"
    outputs: list[Path] = []
    with _open_pdf(src) as doc:
        stem = safe_stem(src.name)
        for page in doc:
            pix = page.get_pixmap(dpi=dpi)
            dest = out_dir / f"{stem}-hal{page.number + 1:03d}.{ext}"
            pix.save(str(dest))
            outputs.append(dest)
    return outputs


def jpg_pdf(inputs: list[Path], opts: dict, out_dir: Path) -> list[Path]:
    from PIL import Image

    out = out_dir / "gambar.pdf"
    prepared: list[Path] = []
    tmp = Path(tempfile.mkdtemp(prefix="img-"))
    try:
        for i, src in enumerate(inputs):
            try:
                im = Image.open(src)
                if im.mode in ("RGBA", "P", "LA"):
                    bg = Image.new("RGB", im.size, (255, 255, 255))
                    im = im.convert("RGBA")
                    bg.paste(im, mask=im.split()[-1])
                    im = bg
                else:
                    im = im.convert("RGB")
                dest = tmp / f"p{i:03d}.jpg"
                im.save(dest, "JPEG", quality=90)
                prepared.append(dest)
            except Exception as exc:
                raise OpError(f"Gambar '{src.name}' tidak bisa dibaca.") from exc
        doc = pymupdf.open()
        for path in prepared:
            img_doc = pymupdf.open(str(path))
            rect = img_doc[0].rect
            page = doc.new_page(width=rect.width, height=rect.height)
            page.insert_image(rect, filename=str(path))
            img_doc.close()
        doc.save(str(out), garbage=4, deflate=True)
        doc.close()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return [out]


def _local_only_fetcher():
    """A fetcher that can only read local files and inline data.

    Without this, WeasyPrint follows ``<img src>`` and ``<link href>`` to any
    address it can reach, which turns an upload into a server-side request
    forger: a file could probe the cloud metadata endpoint (169.254.169.254)
    or any internal service and read the response back through the rendered
    PDF. Restricting the allowed protocols to ``file`` and ``data`` refuses
    http/https/ftp before any connection is made, while a refused resource is
    only a warning: the rest of the document still renders.
    """
    from weasyprint import URLFetcher

    return URLFetcher(allowed_protocols=("file", "data"))


def html_pdf(inputs: list[Path], opts: dict, out_dir: Path) -> list[Path]:
    src = inputs[0]
    out = out_dir / f"{safe_stem(src.name)}.pdf"
    try:
        from weasyprint import HTML

        HTML(filename=str(src), url_fetcher=_local_only_fetcher()).write_pdf(str(out))
        return [out]
    except Exception as exc:
        raise OpError("Gagal merender HTML ke PDF.") from exc


def ocr(inputs: list[Path], opts: dict, out_dir: Path) -> list[Path]:
    import ocrmypdf

    from .. import progress

    src = inputs[0]
    out = out_dir / f"{safe_stem(src.name)}-ocr.pdf"
    with _open_pdf(src):
        pass
    lang = str(opts.get("lang", "ind"))
    if lang not in {"ind", "eng", "ind+eng"}:
        lang = "ind"
    progress.emit(0, None, "Menyiapkan OCR")
    try:
        ocrmypdf.ocr(
            str(src),
            str(out),
            language=lang,
            skip_text=True,
            optimize=1,
            progress_bar=True,
            output_type="pdf",
            # Relays ocrmypdf's own progress bar into this job's progress file.
            plugins=["app.ocr_plugin"],
        )
    except Exception as exc:
        raise OpError(f"OCR gagal: {exc}") from exc
    return [out]
